import io
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from app.main import create_app

app = create_app()
client = TestClient(app)

# 登录
r = client.post("/auth/register", json={"username": "p1user", "password": "p1pass123"})
if r.status_code == 409:
    r = client.post("/auth/login", json={"username": "p1user", "password": "p1pass123"})
assert r.status_code == 200, r.text
token = r.json()["token"]
H = {"Authorization": f"Bearer {token}"}

# 1) DOCX 解析 + 上传
print("==> [1] DOCX 上传解析")
from docx import Document
doc = Document()
doc.add_heading("员工福利制度", 0)
doc.add_paragraph("补充福利：每月团建活动，每年一次体检，生日发放 200 元礼品卡。")
buf = io.BytesIO()
doc.save(buf)
buf.seek(0)
r = client.post("/upload", files={"file": ("员工福利制度.docx", buf.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, headers=H)
print("docx upload:", r.status_code, r.json())
assert r.status_code == 200, r.text

# 2) 多轮对话（conversation_id 延续）
print("==> [2] 多轮对话")
r1 = client.post("/ask", json={"question": "年假有几天？"}, headers=H)
d1 = r1.json()
cid = d1["conversation_id"]
print("round1 mode:", d1["mode"])
r2 = client.post("/ask", json={"question": "那出差住宿标准呢？", "conversation_id": cid}, headers=H)
d2 = r2.json()
print("round2 mode:", d2["mode"], "| same cid:", d2["conversation_id"] == cid)
assert d2["conversation_id"] == cid

# 3) 评测历史（跑一次评测 → 查历史）
print("==> [3] 评测历史")
ts = Path("D:/Projects/smart-kb-agent/data/evaluation/test_set_smart.json")
r3 = client.post("/evaluation/run", json={"test_set_path": str(ts), "top_k": 3}, headers=H)
assert r3.status_code == 200, r3.text
r4 = client.get("/evaluation/runs", headers=H)
runs = r4.json()
print("runs count:", len(runs))
assert len(runs) >= 1, "应有评测历史"

# 4) 前端可访问
print("==> [4] 前端页")
r5 = client.get("/")
print("index status:", r5.status_code)
assert r5.status_code == 200

print("==> P1 验证通过 ==")