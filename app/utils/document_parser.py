"""文档解析：PDF / Markdown / TXT / DOCX → 纯文本。

二进制格式（PDF / DOCX）解析失败时统一转成 `AppError`（400），
而不是把底层异常抛到接口层变成无信息的 500 —— 用户上传损坏文件是可预期的输入。
"""
from __future__ import annotations

from pathlib import Path

from app.utils.exceptions import AppError, UnsupportedFileTypeError

_TEXT_EXTS = {".txt", ".md", ".markdown", ".text"}


def parse_document(file_path: str) -> str:
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext in _TEXT_EXTS:
        return _parse_text(path)
    if ext == ".docx":
        return _parse_docx(path)
    raise UnsupportedFileTypeError(f"不支持的文档类型：{ext}（支持 PDF / Markdown / TXT / DOCX）")


def _parse_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise UnsupportedFileTypeError("缺少 pypdf，请先安装：pip install pypdf") from exc

    try:
        reader = PdfReader(str(path))
        if getattr(reader, "is_encrypted", False):
            raise AppError("PDF 已加密，请先解除密码后再上传", 400)
        pages = [page.extract_text() or "" for page in reader.pages]
    except AppError:
        raise
    except Exception as exc:
        raise AppError(f"PDF 解析失败，文件可能已损坏：{exc}", 400) from exc
    return "\n\n".join(pages)


def _parse_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise UnsupportedFileTypeError("缺少 python-docx，请先安装：pip install python-docx") from exc

    try:
        doc = Document(str(path))
    except AppError:
        raise
    except Exception as exc:
        raise AppError(f"DOCX 解析失败，文件可能已损坏：{exc}", 400) from exc

    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)
